from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

NAVY   = RGBColor(0x1a, 0x2f, 0x6e)
YELLOW = RGBColor(0xf5, 0xc2, 0x00)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAY   = RGBColor(0xAA, 0xAA, 0xBB)
GREEN  = RGBColor(0x16, 0xa3, 0x4a)
ORANGE = RGBColor(0xd9, 0x77, 0x06)
DARK   = RGBColor(0x0f, 0x17, 0x29)
NAVY2  = RGBColor(0x0d, 0x1b, 0x4b)
DKBLUE = RGBColor(0x12, 0x1e, 0x45)

blank = prs.slide_layouts[6]

def rect(slide, l, t, w, h, fill=None, line_color=None):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid() if fill else s.fill.background()
    if fill: s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = Pt(0.75)
    return s

def txt(slide, text, l, t, w, h, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Sarabun"
    return tb

# ══════════════════════════════════════════════
# SLIDE 1 — Cover
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=NAVY)
rect(s, 0, 5.6, 13.33, 1.9, fill=NAVY2)
rect(s, 4.5, 2.0, 4.33, 0.06, fill=YELLOW)  # accent line
txt(s, "SPRINT REVIEW  ·  WuLab Project", 0, 1.1, 13.33, 0.5, size=11, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
txt(s, "ระบบจัดการใบเสนอราคาและคำสั่งซื้อออนไลน์", 0, 1.7, 13.33, 1.0, size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, "บริษัท สยามมาสเตอส์คอนกรีต จำกัด  (SICON)", 0, 2.75, 13.33, 0.6, size=18, color=RGBColor(0xcc,0xd6,0xf0), align=PP_ALIGN.CENTER)
txt(s, "ผู้ผลิตคอนกรีตอัดแรงกว่า 30 ปี  ·  เสาไฟฟ้า  ·  เสาเข็ม  ·  แผ่นพื้น Hollow Core", 0, 3.35, 13.33, 0.5, size=13, color=GRAY, align=PP_ALIGN.CENTER)
txt(s, "Sprint 1: 13 SP     |     Sprint 2: 16 SP     |     8 User Stories     |     Done 8/8 ✓", 0, 6.2, 13.33, 0.5, size=13, color=GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════
# SLIDE 2 — Sprint Planning
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=DARK)
rect(s, 0, 0, 13.33, 1.05, fill=NAVY)
txt(s, "Sprint Planning", 0.4, 0.08, 8, 0.38, size=11, bold=True, color=YELLOW)
txt(s, "User Stories — 2 Sprints", 0.4, 0.42, 10, 0.55, size=22, bold=True, color=WHITE)

sp1 = [
    ("US1","ลูกค้าดูรายการสินค้าคอนกรีตอัดแรงพร้อมสเปค/ราคา","3 SP"),
    ("US2","ลูกค้ากรอกฟอร์ม RFQ 3 แบบตามประเภทลูกค้า → รับรหัส SMC","5 SP"),
    ("US3","ลูกค้าติดตามสถานะคำขอด้วยรหัส SMC","2 SP"),
    ("US4","Admin login + ดูรายการ RFQ + filter + จัดการสถานะ","3 SP"),
]
sp2 = [
    ("US5","Admin สร้างใบเสนอราคา (QT code) + ส่ง link + PDF","5 SP"),
    ("US6","ลูกค้าดูใบเสนอราคา + ทราบสถานะการรับ PO","3 SP"),
    ("US7","Admin บันทึก PO + แนบไฟล์ + auto generate PO code","5 SP"),
    ("US8","Admin อัปเดตสถานะออเดอร์ผ่าน Kanban board","3 SP"),
]

for col, (header, stories) in enumerate([
    ("🚀  Sprint 1 — Customer Flow & Basic Admin  (13 SP)", sp1),
    ("📦  Sprint 2 — Quotation & Order Management  (16 SP)", sp2),
]):
    x = 0.3 + col * 6.55
    rect(s, x, 1.15, 6.25, 5.9, fill=DKBLUE)
    rect(s, x, 1.15, 6.25, 0.42, fill=NAVY)
    txt(s, header, x+0.12, 1.18, 6.0, 0.36, size=11, bold=True, color=YELLOW)
    for i, (us, desc, sp) in enumerate(stories):
        y = 1.68 + i * 1.18
        rect(s, x+0.12, y+0.03, 0.55, 0.36, fill=NAVY)
        txt(s, us, x+0.12, y+0.03, 0.55, 0.36, size=10, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
        txt(s, desc, x+0.73, y+0.03, 4.65, 0.42, size=12, color=WHITE)
        rect(s, x+5.4, y+0.03, 0.72, 0.36, fill=RGBColor(0x1e,0x3a,0x8a))
        txt(s, sp, x+5.4, y+0.03, 0.72, 0.36, size=11, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════
# SLIDE 3 — User Story Status
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=DARK)
rect(s, 0, 0, 13.33, 1.05, fill=NAVY)
txt(s, "User Story Status", 0.4, 0.08, 8, 0.38, size=11, bold=True, color=YELLOW)
txt(s, "To Do  /  Doing  /  Done", 0.4, 0.42, 10, 0.55, size=22, bold=True, color=WHITE)

done_items = [
    "US1 — แสดงสินค้า 3 ประเภท + filter (standard / preorder)",
    "US2 — ฟอร์ม RFQ 3 แบบ + SMC-YYYYMMDD-XXXX auto generate",
    "US3 — Track page SMC code → 3 stage (รอ / ใบเสนอ / PO timeline)",
    "US4 — Admin login + JWT + Dashboard RFQ + badge 'ได้รับ PO'",
    "US5 — สร้างใบเสนอราคา QT code + VAT + shareable link + PDF",
    "US6 — หน้าดูใบเสนอ + status 'รอรับ PO' / 'ได้รับ PO แล้ว'",
    "US7 — บันทึก PO + upload file + PO code + redirect Kanban",
    "US8 — Kanban 4 column + drag & drop + ย้าย status",
]
cols = [
    ("📋  TO DO",    RGBColor(0x22,0x2d,0x44), RGBColor(0x94,0xa3,0xb8), []),
    ("⚡  DOING",    RGBColor(0x2d,0x1a,0x05), ORANGE, []),
    ("✅  DONE (8/8)", RGBColor(0x05,0x20,0x10), GREEN, done_items),
]
for ci, (label, bg, lc, items) in enumerate(cols):
    x = 0.25 + ci * 4.35
    rect(s, x, 1.15, 4.15, 5.95, fill=bg)
    rect(s, x, 1.15, 4.15, 0.42, fill=bg)
    txt(s, label, x+0.12, 1.18, 3.9, 0.36, size=11, bold=True, color=lc)
    if items:
        for j, item in enumerate(items):
            iy = 1.65 + j * 0.63
            rect(s, x+0.1, iy, 3.95, 0.56, fill=RGBColor(0x07,0x30,0x17))
            txt(s, item, x+0.18, iy+0.07, 3.75, 0.45, size=10.5, color=WHITE)
    else:
        txt(s, "— ไม่มีรายการ —", x+0.1, 3.6, 3.9, 0.4, size=12, color=GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════
# SLIDE 4 — Acceptance Criteria
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=DARK)
rect(s, 0, 0, 13.33, 1.05, fill=NAVY)
txt(s, "Acceptance Criteria", 0.4, 0.08, 8, 0.38, size=11, bold=True, color=YELLOW)
txt(s, "เงื่อนไขการยอมรับ", 0.4, 0.42, 10, 0.55, size=22, bold=True, color=WHITE)

ac = [
    ("US2 — ฟอร์มขอใบเสนอราคา", [
        "แบ่ง 3 แบบ: ทั่วไป / เอกชน (บริษัท) / ภาครัฐ (เลขหนังสือ)",
        "เพิ่มสินค้าได้หลายรายการ พร้อมจำนวนและหน่วย",
        "generate SMC-YYYYMMDD-XXXX อัตโนมัติหลัง submit",
        "หน้า confirm แสดงรหัสพร้อม link ไปหน้า track",
    ]),
    ("US3 — ติดตามสถานะ SMC", [
        "กรอก SMC code เดียวติดตามได้ตลอดทุก stage",
        "ยังไม่มีใบเสนอ → แสดงสถานะ + รายการสินค้าที่ขอ",
        "มีใบเสนอ ยังไม่มี PO → วันหมดอายุ (แดงถ้า expired)",
        "มี PO แล้ว → 4-step timeline พร้อมสถานะปัจจุบัน",
    ]),
    ("US5 — สร้างใบเสนอราคา + PDF", [
        "คำนวณ subtotal, VAT 7%, total อัตโนมัติ",
        "generate QT code + link สำหรับส่งลูกค้า",
        "PDF ภาษาไทย ฟอนต์ Sarabun ครบทุกตัวอักษร",
        "สร้างใหม่ (ราคาใหม่) ได้ตราบที่ยังไม่มี PO",
    ]),
    ("US7 — บันทึก PO", [
        "กรอกเลข PO + upload ไฟล์ PDF/JPG/PNG",
        "generate PO-YYYYMMDD-XXXX อัตโนมัติ",
        "สินค้า standard → ข้ามขั้นผลิต ไป 'จัดส่ง' ทันที",
        "redirect ไป Kanban board หลังบันทึกสำเร็จ",
    ]),
]
for ci, (title, items) in enumerate(ac):
    col = ci % 2
    row = ci // 2
    x = 0.3 + col * 6.55
    y = 1.15 + row * 3.05
    rect(s, x, y, 6.3, 2.75, fill=DKBLUE)
    rect(s, x, y, 6.3, 0.4, fill=NAVY)
    txt(s, title, x+0.12, y+0.05, 6.05, 0.33, size=12, bold=True, color=YELLOW)
    for j, item in enumerate(items):
        txt(s, "✓  " + item, x+0.15, y+0.5+j*0.53, 6.0, 0.46, size=11.5, color=WHITE)

# ══════════════════════════════════════════════
# SLIDE 5 — System Demo
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=DARK)
rect(s, 0, 0, 13.33, 1.05, fill=NAVY)
txt(s, "System Demo", 0.4, 0.08, 8, 0.38, size=11, bold=True, color=YELLOW)
txt(s, "User Journey ในระบบ", 0.4, 0.42, 10, 0.55, size=22, bold=True, color=WHITE)

flow = [
    ("US1","ดูสินค้า","/products"),
    ("US2","ส่ง RFQ","/rfq/new"),
    ("US3","Track SMC","/rfq/track"),
    ("US4","Admin จัดการ","/admin"),
    ("US5","ออกใบเสนอ","/admin/rfq/[id]"),
    ("US6","ลูกค้าดู","/quotation"),
    ("US7","รับ PO","admin panel"),
    ("US8","Kanban","/admin/orders"),
]
fw = 1.46
for i, (us, label, url) in enumerate(flow):
    x = 0.22 + i * (fw + 0.15)
    rect(s, x, 1.15, fw, 1.1, fill=NAVY)
    txt(s, us,    x+0.05, 1.17, fw-0.1, 0.28, size=9, bold=True, color=YELLOW)
    txt(s, label, x+0.05, 1.42, fw-0.1, 0.5,  size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, url,   x+0.05, 2.05, fw-0.1, 0.2,  size=8,  color=GRAY, align=PP_ALIGN.CENTER)
    if i < len(flow)-1:
        txt(s, "→", x+fw+0.01, 1.55, 0.13, 0.3, size=12, color=GRAY, align=PP_ALIGN.CENTER)

headers = ["US", "URL / หน้า", "Feature ที่ Demo"]
col_w   = [0.8, 3.2, 8.5]
tx = 0.25
rect(s, tx, 2.5, 12.8, 0.38, fill=NAVY)
cx = tx
for h, cw in zip(headers, col_w):
    txt(s, h, cx+0.1, 2.53, cw, 0.32, size=11, bold=True, color=GRAY)
    cx += cw

rows = [
    ("US1","/products",          "สินค้า 3 ประเภท + badge preorder + ตารางราคา + filter"),
    ("US2","/rfq/new",           "ฟอร์ม 3 tab + validate + submit → รหัส SMC อัตโนมัติ"),
    ("US3","/rfq/track",         "กรอก SMC → เห็น stage: รอ / ใบเสนอ / 4-step PO timeline"),
    ("US4","/admin",             "login + JWT + Dashboard + filter + badge 'ได้รับ PO'"),
    ("US5","/admin/rfq/[id]",    "กรอกราคา → VAT → QT code → copy link → download PDF"),
    ("US6","/quotation?ref=QT-…","ดูใบเสนอ + กล่องสถานะ PO (เหลือง/เขียว)"),
    ("US7","/admin/rfq/[id]",    "upload PO file → PO code → สินค้า standard ข้ามขั้นผลิต"),
    ("US8","/admin/orders",      "Kanban 4 column + drag&drop + ย้าย status"),
]
for ri, row in enumerate(rows):
    ry = 2.95 + ri * 0.52
    bg = RGBColor(0x12,0x1a,0x35) if ri%2==0 else DARK
    rect(s, tx, ry, 12.8, 0.5, fill=bg)
    cx = tx
    for val, cw in zip(row, col_w):
        color = YELLOW if cx == tx else WHITE
        txt(s, val, cx+0.1, ry+0.07, cw-0.1, 0.38, size=11, color=color)
        cx += cw

# ══════════════════════════════════════════════
# SLIDE 6 — Summary
# ══════════════════════════════════════════════
s = prs.slides.add_slide(blank)
rect(s, 0, 0, 13.33, 7.5, fill=NAVY)
rect(s, 0, 5.7, 13.33, 1.8, fill=NAVY2)
txt(s, "สรุปผลการพัฒนา", 0, 0.5, 13.33, 0.6, size=13, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
txt(s, "WuLab — ระบบจัดการใบเสนอราคาออนไลน์ SICON", 0, 1.1, 13.33, 0.8, size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

cards = [("8","User Stories"),("29","Story Points"),("100%","Done ✓")]
for i, (num, label) in enumerate(cards):
    x = 1.5 + i*3.7
    rect(s, x, 2.2, 3.0, 1.95, fill=RGBColor(0x0d,0x1f,0x55))
    txt(s, num,   x, 2.3, 3.0, 1.1, size=52, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
    txt(s, label, x, 3.3, 3.0, 0.5, size=16, color=WHITE,  align=PP_ALIGN.CENTER)

txt(s, "Sprint 1 (US1–US4): 13 SP   |   Sprint 2 (US5–US8): 16 SP",
    0, 4.45, 13.33, 0.5, size=14, color=GRAY, align=PP_ALIGN.CENTER)
txt(s, "Next.js 16  ·  Prisma 7  ·  SQLite  ·  TypeScript  ·  Tailwind CSS",
    0, 5.0, 13.33, 0.45, size=12, color=GRAY, align=PP_ALIGN.CENTER)
txt(s, "บริษัท สยามมาสเตอส์คอนกรีต จำกัด  (SICON)  ·  ก่อตั้ง พ.ศ. 2537",
    0, 6.35, 13.33, 0.45, size=12, color=GRAY, align=PP_ALIGN.CENTER)

out = r"C:\Users\Win11\Desktop\CLAUDE MAX\Project\WuLab\docs\SICON-Sprint-Presentation.pptx"
prs.save(out)
print("Saved:", out)
